"""Porsche sales sanitization agent.

Reads a raw Porsche sales workbook, applies the rules in schema.md and writes a
new workbook where each sanitized column sits right after its source column.

Usage:
    python sanitize_porsche.py INPUT.xlsx [OUTPUT.xlsx]
"""
import re
import sys
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

INVALID = "INVALID"

# ---------------------------------------------------------------- helpers

def text(v):
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def title_case(s):
    # Capitalize each word, keeping punctuation like "St." intact.
    return " ".join(w[:1].upper() + w[1:].lower() for w in text(s).split(" ") if w)


UNITS = {
    "zero": 0, "one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6,
    "seven": 7, "eight": 8, "nine": 9, "ten": 10, "eleven": 11, "twelve": 12,
    "thirteen": 13, "fourteen": 14, "fifteen": 15, "sixteen": 16,
    "seventeen": 17, "eighteen": 18, "nineteen": 19,
}
TENS = {
    "twenty": 20, "thirty": 30, "forty": 40, "fifty": 50, "sixty": 60,
    "seventy": 70, "eighty": 80, "ninety": 90,
}
SCALES = {"thousand": 1_000, "million": 1_000_000}


def words_to_number(s):
    """'two hundred thousand' -> 200000. Returns None if a word is unknown."""
    tokens = [t for t in re.split(r"[\s-]+", s.lower()) if t and t != "and"]
    if not tokens:
        return None
    total = current = 0
    for t in tokens:
        if t in UNITS:
            current += UNITS[t]
        elif t in TENS:
            current += TENS[t]
        elif t == "hundred":
            current = (current or 1) * 100
        elif t in SCALES:
            total += (current or 1) * SCALES[t]
            current = 0
        else:
            return None
    return total + current


def parse_amount(s):
    """Parse '1,234.56', '1.234,56', '119.900' or '9.5' into a Decimal."""
    s = s.strip()
    if "," in s and "." in s:
        # The right-most separator is the decimal one.
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", "") if re.fullmatch(r"\d{1,3}(,\d{3})+", s) else s.replace(",", ".")
    elif "." in s and re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")  # European thousands separator, e.g. 119.900
    return Decimal(s)

# ---------------------------------------------------------------- dates

MONTHS = {m.lower(): i for i, m in enumerate(
    ["January", "February", "March", "April", "May", "June", "July",
     "August", "September", "October", "November", "December"], 1)}
MONTHS.update({k[:3]: v for k, v in list(MONTHS.items())})
MONTHS["sept"] = 9


def make_date(y, m, d):
    try:
        return date(int(y), int(m), int(d)).isoformat()
    except ValueError:
        return INVALID


def sanitize_date(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = text(v)
    if m := re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s):
        return make_date(*m.groups())
    if m := re.fullmatch(r"(\d{1,2})[-/](\d{1,2})[-/](\d{4}|\d{2})", s):
        mo, d, y = m.groups()
        return make_date(int(y) + 2000 if len(y) == 2 else y, mo, d)
    if m := re.fullmatch(r"([A-Za-z]+)\.?\s+(\d{1,2})(?:st|nd|rd|th)?,?\s+(\d{4})", s, re.I):
        mon, d, y = m.groups()
        if mon.lower() in MONTHS:
            return make_date(y, MONTHS[mon.lower()], d)
    return INVALID

# ---------------------------------------------------------------- models

CANONICAL_MODELS = [
    "911 Carrera", "911 Carrera S", "911 Carrera GTS", "911 Turbo", "911 Turbo S",
    "911 GT3", "911 GT3 RS", "911 Dakar", "911 Targa 4", "911 Targa 4S",
    "718 Cayman", "718 Cayman S", "718 Cayman GT4 RS", "718 Boxster",
    "718 Boxster GTS", "718 Spyder RS", "Cayenne", "Cayenne S", "Cayenne Coupe",
    "Cayenne E-Hybrid", "Cayenne Turbo", "Cayenne Turbo GT", "Macan", "Macan S",
    "Macan T", "Macan GTS", "Macan Electric", "Panamera", "Panamera 4",
    "Panamera 4S", "Panamera Turbo", "Panamera Turbo S", "Panamera 4 E-Hybrid",
    "Taycan", "Taycan 4S", "Taycan GTS", "Taycan Turbo", "Taycan Turbo S",
    "Taycan Cross Turismo",
]


def model_key(s):
    return re.sub(r"[\s\-_]+", " ", s.lower()).replace("e hybrid", "e-hybrid").strip()


MODEL_LOOKUP = {model_key(m): m for m in CANONICAL_MODELS}
# Tokens that keep their upper-case form when title-casing an unknown model.
UPPER_TOKENS = {"gt", "gt2", "gt3", "gt4", "gts", "rs", "4s", "s", "t"}


def sanitize_model(v):
    s = text(v)
    if not s:
        return INVALID
    if (hit := MODEL_LOOKUP.get(model_key(s))):
        return hit
    return " ".join(w.upper() if w.lower() in UPPER_TOKENS else title_case(w) for w in s.split(" "))

# ---------------------------------------------------------------- model year

def sanitize_year(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        year = int(v) if float(v).is_integer() else None
    else:
        s = text(v).lower()
        if re.fullmatch(r"\d{2}[\s\-]?\d{2}", s):
            year = int(re.sub(r"\D", "", s))
        elif re.fullmatch(r"\d{4}(\.0)?", s):
            year = int(float(s))
        else:
            year = words_year(s)
    return str(year) if year is not None and 1990 <= year <= 2035 else INVALID


def words_year(s):
    # "twenty twenty four" -> 20 + 24; "two thousand twenty one" -> 2021
    if "thousand" in s:
        return words_to_number(s)
    parts = s.split()
    if len(parts) >= 2:
        hi, lo = words_to_number(parts[0]), words_to_number(" ".join(parts[1:]))
        if hi is not None and lo is not None and lo < 100:
            return hi * 100 + lo
    return None

# ---------------------------------------------------------------- price

def sanitize_price(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        amount = Decimal(str(v))
    else:
        s = text(v).lower()
        s = re.sub(r"\b(usd|us\$|dollars?)\b|\$", " ", s).strip()
        try:
            if m := re.fullmatch(r"([\d.,]+)\s*k", s):
                amount = parse_amount(m.group(1)) * 1000
            elif re.fullmatch(r"[\d.,]+", s):
                amount = parse_amount(s)
            elif (n := words_to_number(s)) is not None:
                amount = Decimal(n)
            else:
                return INVALID
        except ArithmeticError:
            return INVALID
    return str(amount.quantize(Decimal("0.01"), ROUND_HALF_UP))

# ---------------------------------------------------------------- mileage

KM_TO_MILES = Decimal("0.621371")


def sanitize_mileage(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        miles = Decimal(str(v))
    else:
        s = text(v).lower()
        if s in {"new", "new car", "zero", "zero miles", "brand new"}:
            return "0"
        is_km = bool(re.search(r"\b(km|kms|kilometers?|kilometres?)\b", s))
        s = re.sub(r"\b(km|kms|kilometers?|kilometres?|miles?|mi)\b\.?|:", " ", s)
        s = re.sub(r"(\d)(mi|km)\b", r"\1", s).strip()
        try:
            if re.fullmatch(r"[\d.,]+", s):
                miles = parse_amount(s)
            elif (n := words_to_number(s)) is not None:
                miles = Decimal(n)
            else:
                return INVALID
        except ArithmeticError:
            return INVALID
        if is_km:
            miles *= KM_TO_MILES
    return str(int(miles.quantize(Decimal("1"), ROUND_HALF_UP)))

# ---------------------------------------------------------------- payment

PAYMENT_RULES = [
    (r"crypto|bitcoin|btc", "Crypto Payment"),
    (r"\bach\b", "ACH Payment"),
    (r"debit", "Debit Card"),
    (r"credit", "Credit Card"),
    (r"wire", "Wire Transfer"),
    (r"bank ?transfer", "Bank Transfer"),
    (r"financ|loan", "Financing"),
    (r"leas", "Lease"),
    (r"cash", "Cash"),
]


def sanitize_payment(v):
    s = text(v)
    if not s:
        return INVALID
    # Split CamelCase and separators: "CreditCard" / "bank-transfer" -> words.
    norm = re.sub(r"([a-z])([A-Z])", r"\1 \2", s)
    norm = re.sub(r"[\-_]+", " ", norm).lower()
    for pattern, label in PAYMENT_RULES:
        if re.search(pattern, norm):
            return label
    return title_case(norm)

# ---------------------------------------------------------------- city / state

def sanitize_city(v):
    s = text(v)
    return title_case(s) if s else INVALID


STATES = {
    "AL": "Alabama", "AK": "Alaska", "AZ": "Arizona", "AR": "Arkansas",
    "CA": "California", "CO": "Colorado", "CT": "Connecticut", "DE": "Delaware",
    "DC": "District of Columbia", "FL": "Florida", "GA": "Georgia", "HI": "Hawaii",
    "ID": "Idaho", "IL": "Illinois", "IN": "Indiana", "IA": "Iowa", "KS": "Kansas",
    "KY": "Kentucky", "LA": "Louisiana", "ME": "Maine", "MD": "Maryland",
    "MA": "Massachusetts", "MI": "Michigan", "MN": "Minnesota", "MS": "Mississippi",
    "MO": "Missouri", "MT": "Montana", "NE": "Nebraska", "NV": "Nevada",
    "NH": "New Hampshire", "NJ": "New Jersey", "NM": "New Mexico", "NY": "New York",
    "NC": "North Carolina", "ND": "North Dakota", "OH": "Ohio", "OK": "Oklahoma",
    "OR": "Oregon", "PA": "Pennsylvania", "RI": "Rhode Island", "SC": "South Carolina",
    "SD": "South Dakota", "TN": "Tennessee", "TX": "Texas", "UT": "Utah",
    "VT": "Vermont", "VA": "Virginia", "WA": "Washington", "WV": "West Virginia",
    "WI": "Wisconsin", "WY": "Wyoming",
}
STATE_LOOKUP = {**{k.lower(): k for k in STATES}, **{v.lower(): k for k, v in STATES.items()}}


def sanitize_state(v):
    s = re.sub(r"[.\s]+", " ", text(v)).strip().lower()
    return STATE_LOOKUP.get(s, INVALID)

# ---------------------------------------------------------------- delivery

DELIVERY_LABELS = [
    "Delivered", "Pending", "In Transit", "Cancelled", "Awaiting Delivery",
    "Awaiting Pickup", "Pending Approval", "Pending Review", "Shipped",
    "Awaiting Review",
]
DELIVERY_LOOKUP = {d.lower(): d for d in DELIVERY_LABELS}
DELIVERY_LOOKUP.update({"deliverd": "Delivered", "canceled": "Cancelled", "intransit": "In Transit"})


def sanitize_delivery(v):
    s = re.sub(r"[^a-z ]+", " ", text(v).lower())
    s = re.sub(r"\s+", " ", s).strip()
    if not s:
        return INVALID
    return DELIVERY_LOOKUP.get(s, title_case(s))

# ---------------------------------------------------------------- pipeline

RULES = {  # source column -> (output column, sanitizer)
    "sale_date": ("SaleDateSanitized", sanitize_date),
    "porsche_model": ("PorscheModelSanitized", sanitize_model),
    "model_year": ("ModelYearSanitized", sanitize_year),
    "sale_price": ("SalesPriceSanitized", sanitize_price),
    "vehicle_mileage": ("VehicleMileageSanitized", sanitize_mileage),
    "payment_method": ("PayMethodSanitized", sanitize_payment),
    "city": ("CitySanitized", sanitize_city),
    "state": ("StateSanitized", sanitize_state),
    "delivery_status": ("DeliveryStatusSanitized", sanitize_delivery),
}
REQUIRED = ["sale_id", "sale_date", "customer_name", "porsche_model", "model_year",
            "sale_price", "vehicle_mileage", "payment_method", "city", "state",
            "salesperson", "delivery_status"]


def quality_checks(header, rows):
    """Return a list of (check, passed, detail) per the schema's quality checks."""
    idx = {h: i for i, h in enumerate(header)}
    col = lambda name: [r[idx[name]] for r in rows]
    checks = [
        ("No original columns were removed", all(c in idx for c in REQUIRED), ""),
        ("Sanitized columns follow their source",
         all(idx[out] == idx[src] + 1 for src, (out, _) in RULES.items()), ""),
        ("Dates are ISO or INVALID",
         all(v == INVALID or re.fullmatch(r"\d{4}-\d{2}-\d{2}", v) for v in col("SaleDateSanitized")), ""),
        ("States are USPS codes or INVALID",
         all(v == INVALID or v in STATES for v in col("StateSanitized")), ""),
        ("Mileage is integer miles",
         all(v == INVALID or re.fullmatch(r"\d+", v) for v in col("VehicleMileageSanitized")), ""),
        ("Price has two decimals",
         all(v == INVALID or re.fullmatch(r"\d+\.\d{2}", v) for v in col("SalesPriceSanitized")), ""),
        ("No blank sanitized values",
         all(text(r[idx[out]]) for r in rows for out, _ in RULES.values()), ""),
    ]
    return checks


def run(input_path, output_path):
    wb_in = load_workbook(input_path)
    ws_in = wb_in.worksheets[0]
    raw = list(ws_in.iter_rows(values_only=True))
    src_header = [text(h) for h in raw[0]]
    missing = [c for c in REQUIRED if c not in src_header]
    if missing:
        raise SystemExit(f"Missing required columns: {missing}")

    header, plan = [], []
    for i, name in enumerate(src_header):
        header.append(name)
        plan.append((i, None))
        if name in RULES:
            header.append(RULES[name][0])
            plan.append((i, RULES[name][1]))

    rows = []
    for r in raw[1:]:
        if all(v is None or text(v) == "" for v in r):
            continue
        rows.append([fn(r[i]) if fn else r[i] for i, fn in plan])

    # Write output workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sanitized"
    ws.append(header)
    for row in rows:
        ws.append([v.strftime("%Y-%m-%d") if isinstance(v, (datetime, date)) else v for v in row])
    base, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    hdr_fill = PatternFill("solid", fgColor="D9E1F2")
    san_fill = PatternFill("solid", fgColor="E2EFDA")
    bad = Font(name="Arial", size=10, color="C00000", bold=True)
    sanitized_cols = {i + 1 for i, h in enumerate(header) if h.endswith("Sanitized")}
    for row in ws.iter_rows():
        for c in row:
            c.font = base
            if c.row == 1:
                c.font, c.fill = bold, (san_fill if c.column in sanitized_cols else hdr_fill)
            elif c.value == INVALID:
                c.font = bad
    for i, h in enumerate(header, 1):
        width = max(len(h), *(len(text(r[i - 1])) for r in rows)) + 2
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(width, 40)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    # Report sheet: quality checks + INVALID counts
    checks = quality_checks(header, rows)
    rep = wb.create_sheet("Quality Report")
    rep.append(["Check", "Result"])
    for name, ok, _ in checks:
        rep.append([name, "PASS" if ok else "FAIL"])
    rep.append([])
    rep.append(["Sanitized column", "INVALID count"])
    invalid_counts = {}
    for out, _ in RULES.values():
        n = sum(1 for r in rows if r[header.index(out)] == INVALID)
        invalid_counts[out] = n
        rep.append([out, n])
    rep.append([])
    rep.append(["Rows processed", len(rows)])
    rep.append(["Source file", Path(input_path).name])
    for row in rep.iter_rows():
        for c in row:
            c.font = bold if c.row == 1 or c.value in ("Sanitized column",) else base
    rep.column_dimensions["A"].width = 42
    rep.column_dimensions["B"].width = 18

    wb.save(output_path)
    return rows, checks, invalid_counts


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_name(src.stem + "_sanitized.xlsx")
    rows, checks, invalid = run(src, dst)
    print(f"Rows processed: {len(rows)}")
    for name, ok, _ in checks:
        print(f"[{'PASS' if ok else 'FAIL'}] {name}")
    print("INVALID counts:", ", ".join(f"{k}={v}" for k, v in invalid.items()))
    print(f"Output: {dst}")
    if not all(ok for _, ok, _ in checks):
        sys.exit(1)


if __name__ == "__main__":
    main()
